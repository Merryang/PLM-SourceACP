"""Build an audit report for exact sequence overlaps in benchmark splits."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(".")
BRANCH = ROOT / "meta_acp" / "dataset_branches" / "benchmark_prott5_mscrnet_cdhit09"
OUT_DIR = ROOT / "meta_acp" / "reports" / "overlap_audit"


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width: float) -> None:
    cell.width = Inches(width)
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_table(table, widths: list[float]) -> None:
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[col_idx])
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
            if row_idx == 0:
                shade_cell(cell, "E8EEF5")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def source_detail(raw: pd.DataFrame, sequence: str, source_file: str) -> dict:
    sub = raw[(raw["sequence"] == sequence) & (raw["source_file"] == source_file)].copy()
    labels = ";".join(map(str, sorted(sub["binary_label"].astype(int).unique())))
    raw_labels = ";".join(sorted(sub["raw_label"].astype(str).unique()))
    ids = ";".join(sub["source_id"].astype(str).tolist())
    return {"labels": labels, "raw_labels": raw_labels, "source_ids": ids, "n": int(len(sub))}


def build_duplicate_tables() -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    overlap = pd.read_csv(BRANCH / "exact_overlap_sequences.csv")
    raw = pd.read_csv(ROOT / "raw_records_label_audit.csv")

    train_test2 = overlap[
        ((overlap["left_source"] == "Test2.fasta") & (overlap["right_source"] == "Train.fasta"))
        | ((overlap["left_source"] == "Train.fasta") & (overlap["right_source"] == "Test2.fasta"))
    ].copy()
    train_test2 = train_test2.sort_values(["binary_labels", "length", "sequence"]).reset_index(drop=True)

    records = []
    for idx, row in train_test2.iterrows():
        seq = row["sequence"]
        train = source_detail(raw, seq, "Train.fasta")
        test2 = source_detail(raw, seq, "Test2.fasta")
        consistent = train["labels"] == test2["labels"] and row["binary_labels"] in {"0", "1"}
        if row["binary_labels"] == "0;1":
            note = "Label conflict between Train and Test2"
        elif consistent:
            note = "Exact overlap with same label"
        else:
            note = "Exact overlap with label discrepancy"
        records.append(
            {
                "No": idx + 1,
                "sequence": seq,
                "length": int(row["length"]),
                "combined_labels": row["binary_labels"],
                "train_label": train["labels"],
                "train_raw_label": train["raw_labels"],
                "train_source_ids": train["source_ids"],
                "test2_label": test2["labels"],
                "test2_raw_label": test2["raw_labels"],
                "test2_source_ids": test2["source_ids"],
                "label_consistency": "consistent" if consistent else "conflict",
                "note": note,
            }
        )
    detailed = pd.DataFrame(records)

    pair_summary = (
        overlap.groupby(["left_source", "right_source", "binary_labels"])
        .size()
        .reset_index(name="n_sequences")
        .sort_values(["left_source", "right_source", "binary_labels"])
    )
    summary = {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "train_test2_overlap": int(len(detailed)),
        "train_test2_non_acp": int((detailed["combined_labels"] == "0").sum()),
        "train_test2_acp": int((detailed["combined_labels"] == "1").sum()),
        "train_test2_label_conflict": int((detailed["combined_labels"] == "0;1").sum()),
        "test1_test2_overlap": int(
            len(overlap[(overlap["left_source"] == "Test1.fasta") & (overlap["right_source"] == "Test2.fasta")])
        ),
        "train_test1_overlap": int(
            len(overlap[(overlap["left_source"] == "Test1.fasta") & (overlap["right_source"] == "Train.fasta")])
        ),
        "length_min": int(detailed["length"].min()),
        "length_median": float(detailed["length"].median()),
        "length_mean": float(detailed["length"].mean()),
        "length_max": int(detailed["length"].max()),
    }
    return detailed, pair_summary, summary


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(46, 116, 181)


def add_summary_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    headers = ["Item", "Value", "Interpretation"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, size=9)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text, size=8)
    style_table(table, [2.2, 1.2, 5.8])


def add_dataframe_table(doc: Document, df: pd.DataFrame, columns: list[str], widths: list[float], max_rows: int | None = None) -> None:
    part = df if max_rows is None else df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, col in zip(table.rows[0].cells, columns):
        set_cell_text(cell, col, bold=True, size=7)
    for _, row in part.iterrows():
        cells = table.add_row().cells
        for cell, col in zip(cells, columns):
            size = 6 if col == "sequence" else 7
            set_cell_text(cell, row[col], size=size)
    style_table(table, widths)


def build_docx(detailed: pd.DataFrame, pair_summary: pd.DataFrame, summary: dict, out_docx: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    for side in ["top_margin", "bottom_margin", "left_margin", "right_margin"]:
        setattr(section, side, Inches(0.7))

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title_run = title.add_run("Train-Test2 Exact Sequence Overlap Audit")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(11, 37, 69)
    doc.add_paragraph("ProtT5-MSCRNet-compatible benchmark branch; generated for Results and Discussion writing.")
    doc.add_paragraph(f"Generated: {summary['generated_at']}")

    add_heading(doc, "Key Findings", 1)
    doc.add_paragraph(
        "The formal ProtT5-MSCRNet benchmark files contain exact sequence overlap between Train.fasta and Test2.fasta. "
        "This is important because Test2 is described as an independent evaluation set; exact overlap with the training "
        "file may inflate apparent generalization performance, especially when duplicated sequences retain the same label."
    )
    add_summary_table(
        doc,
        [
            ["Train-Test2 exact duplicates", str(summary["train_test2_overlap"]), "Exact sequence-level overlap between the training file and Test2."],
            ["non-ACP duplicates", str(summary["train_test2_non_acp"]), "Most duplicated Train-Test2 sequences are non-ACP."],
            ["ACP duplicates", str(summary["train_test2_acp"]), "Duplicated positive examples can increase Test2 sensitivity."],
            ["label-conflict duplicate", str(summary["train_test2_label_conflict"]), "One exact sequence appears with both labels and should be discussed or excluded in strict audits."],
            ["Test1-Test2 exact duplicates", str(summary["test1_test2_overlap"]), "Test1 and Test2 also overlap, affecting independence between evaluation sets."],
            ["Train-Test1 exact duplicates", str(summary["train_test1_overlap"]), "No exact overlap was detected between Train and Test1 in this audit."],
            [
                "Length range",
                f"{summary['length_min']}-{summary['length_max']} aa",
                f"Median length {summary['length_median']:.1f} aa; mean length {summary['length_mean']:.1f} aa.",
            ],
        ],
    )

    add_heading(doc, "Pairwise Overlap Summary", 1)
    add_dataframe_table(
        doc,
        pair_summary,
        ["left_source", "right_source", "binary_labels", "n_sequences"],
        [1.5, 1.5, 1.2, 1.2],
    )

    add_heading(doc, "Interpretation for the Manuscript", 1)
    doc.add_paragraph(
        "For fair comparison with ProtT5-MSCRNet, the official Train/Test1/Test2 counts should be preserved. "
        "However, the exact Train-Test2 overlap means that high Test2 performance should not be interpreted as a fully "
        "homology-independent generalization result. In our main integrated benchmark, global CD-HIT 0.8 redundancy "
        "reduction is therefore a stricter evaluation design. The ProtT5-MSCRNet-compatible branch should be reported "
        "as a literature-compatible comparison, while the CD-HIT 0.8 integrated branch should remain the primary "
        "provenance-aware generalization benchmark."
    )
    doc.add_paragraph(
        "Suggested Discussion wording: 'An audit of the formal ProtT5-MSCRNet benchmark revealed exact sequence overlap "
        "between Train.fasta and Test2.fasta. Therefore, although we include a literature-compatible comparison using "
        "the original partitions, we regard the globally deduplicated integrated benchmark as the stricter assessment "
        "of cross-source generalization.'"
    )

    conflict = detailed[detailed["combined_labels"] == "0;1"]
    add_heading(doc, "Label-Conflict Duplicate", 1)
    if len(conflict):
        add_dataframe_table(
            doc,
            conflict,
            ["No", "length", "combined_labels", "sequence", "train_source_ids", "test2_source_ids", "note"],
            [0.35, 0.45, 0.8, 4.0, 1.5, 1.5, 1.9],
        )
    else:
        doc.add_paragraph("No Train-Test2 exact duplicate had conflicting labels.")

    add_heading(doc, "Complete Train-Test2 Duplicate Sequence List", 1)
    doc.add_paragraph(
        "The table below lists all exact Train-Test2 duplicate sequences. The same information is also exported as CSV "
        "for filtering and supplementary-table preparation."
    )
    add_dataframe_table(
        doc,
        detailed,
        ["No", "combined_labels", "length", "sequence", "train_source_ids", "test2_source_ids", "note"],
        [0.35, 0.7, 0.45, 4.5, 1.35, 1.35, 1.65],
    )

    doc.save(out_docx)


def build_html(detailed: pd.DataFrame, pair_summary: pd.DataFrame, summary: dict, out_html: Path) -> None:
    html = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>Train-Test2 Overlap Audit</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;line-height:1.55;margin:28px;color:#15202b}}
h1{{color:#0b2545}} h2{{color:#2e74b5;margin-top:28px}}
table{{border-collapse:collapse;width:100%;font-size:12px}} th,td{{border:1px solid #d8e0e8;padding:6px;vertical-align:top}}
th{{background:#e8eef5}} code{{font-family:Consolas,monospace;font-size:11px;word-break:break-all}}
.callout{{background:#f4f6f9;border-left:4px solid #2e74b5;padding:10px 14px;margin:12px 0}}
</style></head><body>
<h1>Train-Test2 Exact Sequence Overlap Audit</h1>
<p>Generated: {summary['generated_at']}</p>
<div class="callout">Train-Test2 exact duplicates: <b>{summary['train_test2_overlap']}</b>;
non-ACP: <b>{summary['train_test2_non_acp']}</b>; ACP: <b>{summary['train_test2_acp']}</b>;
label conflict: <b>{summary['train_test2_label_conflict']}</b>.</div>
<h2>Pairwise overlap summary</h2>
{pair_summary.to_html(index=False, escape=False)}
<h2>Complete Train-Test2 duplicate sequence list</h2>
{detailed.to_html(index=False, escape=False)}
</body></html>"""
    out_html.write_text(html, encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    detailed, pair_summary, summary = build_duplicate_tables()
    detailed.to_csv(OUT_DIR / "train_test2_exact_duplicates_detailed.csv", index=False)
    pair_summary.to_csv(OUT_DIR / "pairwise_exact_overlap_summary.csv", index=False)
    (OUT_DIR / "train_test2_overlap_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    build_html(detailed, pair_summary, summary, OUT_DIR / "train_test2_exact_overlap_report.html")
    build_docx(detailed, pair_summary, summary, OUT_DIR / "Train_Test2_exact_overlap_audit_report.docx")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
